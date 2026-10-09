"""Generate the human decision workbook for blocked Firebase migration waves."""

from __future__ import annotations

import argparse
import json
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation


def safe(value: Any) -> Any:
    if isinstance(value, (dict, list)):
        value = json.dumps(value, ensure_ascii=False)
    if isinstance(value, str) and value.startswith(("=", "+", "-", "@")):
        return "'" + value
    return value


def add_sheet(workbook: Workbook, title: str, headers: list[str], rows: list[list[Any]]) -> Any:
    sheet = workbook.create_sheet(title)
    sheet.append(headers)
    for row in rows:
        sheet.append([safe(value) for value in row])
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for column in sheet.columns:
        letter = column[0].column_letter
        width = min(60, max(12, max(len(str(cell.value or "")) for cell in column) + 2))
        sheet.column_dimensions[letter].width = width
    return sheet


def add_list_validation(sheet, column_letter: str, values: str, row_count: int) -> None:
    validation = DataValidation(type="list", formula1=f'"{values}"', allow_blank=True)
    validation.error = "Selecione uma opcao da lista."
    validation.errorTitle = "Valor invalido"
    sheet.add_data_validation(validation)
    validation.add(f"{column_letter}2:{column_letter}{row_count + 1}")


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--source", required=True)
    parser.add_argument("--reports", default="reports/firebase_migration")
    parser.add_argument("--output", required=True)
    args = parser.parse_args()
    source = json.loads(Path(args.source).read_text(encoding="utf-8"))
    reports = Path(args.reports)
    supplier_map = json.loads((reports / "SUPPLIER_ID_MAP.json").read_text(encoding="utf-8"))["entries"]
    material_map = json.loads((reports / "MATERIAL_ID_MAP.json").read_text(encoding="utf-8"))["entries"]
    sku_map = json.loads((reports / "SKU_ID_MAP.json").read_text(encoding="utf-8"))["entries"]

    workbook = Workbook()
    instructions = workbook.active
    instructions.title = "INSTRUCOES"
    instruction_rows = [
        ["Objetivo", "Registrar somente decisoes humanas; esta planilha nao aplica dados automaticamente."],
        ["Fornecedores", "Corrigir, unificar ou ignorar os 20 casos bloqueados."],
        ["Materiais", "Confirmar dominio MP, EP, ES, RT ou FRAGRANCIA e decidir criar/mapear/ignorar."],
        ["SKUs", "Resolver por alias, manter historico autorizado ou ignorar. Nao criar SKU automaticamente."],
        ["Seguranca", "Nao altere as colunas LEGACY_ID. Preencha apenas as colunas DECISAO/APROVACAO."],
    ]
    for row in instruction_rows:
        instructions.append(row)
    instructions.column_dimensions["A"].width = 22
    instructions.column_dimensions["B"].width = 110
    instructions["A1"].font = Font(bold=True)

    blocked_supplier_rows = []
    for entry in supplier_map:
        if entry["classification"] == "INSERT_CANDIDATE":
            continue
        payload = source.get("fornecedores", {}).get(entry["legacy_id"], {})
        blocked_supplier_rows.append([
            entry["legacy_id"], entry["classification"], entry["reason"],
            payload.get("razaoSocial", ""), payload.get("nomeFantasia", ""), payload.get("cnpj", ""),
            payload.get("cidade", ""), payload.get("uf", ""), payload.get("contatoNome", ""),
            payload.get("contatoTelefone", ""), payload.get("contatoEmail", ""),
            "", "", "", "",
        ])
    supplier_sheet = add_sheet(workbook, "FORNECEDORES_BLOQUEADOS", [
        "LEGACY_ID", "CLASSIFICACAO", "MOTIVO", "RAZAO_SOCIAL", "NOME_FANTASIA", "CNPJ_ORIGINAL",
        "CIDADE", "UF", "CONTATO", "TELEFONE", "EMAIL", "DECISAO", "CNPJ_CORRIGIDO",
        "FORNECEDOR_ALVO_ID", "OBSERVACAO",
    ], blocked_supplier_rows)
    add_list_validation(supplier_sheet, "L", "CORRIGIR,UNIFICAR,IGNORAR", len(blocked_supplier_rows))

    material_rows = []
    for entry in material_map:
        payload = source.get("materiais", {}).get(entry["legacy_id"], {})
        material_rows.append([
            entry["legacy_id"], entry.get("legacy_material_code"), payload.get("mpNome", ""),
            entry.get("legacy_type"), entry.get("target_domain"), entry.get("confidence"), entry.get("reason"),
            payload.get("unidade", ""), payload.get("unidadeCompra", ""), payload.get("controlado", ""),
            payload.get("codigoBarras", ""), payload.get("especificacoesTecnicas", ""),
            "", "", "", "", "",
        ])
    material_sheet = add_sheet(workbook, "MATERIAIS", [
        "LEGACY_ID", "CODIGO_LEGADO", "NOME", "TIPO_LEGADO", "DOMINIO_SUGERIDO_NAO_APROVADO",
        "CONFIANCA", "MOTIVO", "UNIDADE", "UNIDADE_COMPRA", "CONTROLADO", "CODIGO_BARRAS",
        "ESPECIFICACOES", "DECISAO", "DOMINIO_APROVADO", "TARGET_ID", "CODIGO_FINAL", "OBSERVACAO",
    ], material_rows)
    add_list_validation(material_sheet, "M", "CRIAR,MAPEAR,IGNORAR", len(material_rows))
    add_list_validation(material_sheet, "N", "MP,EP,ES,RT,FRAGRANCIA", len(material_rows))

    sku_rows = []
    for entry in sku_map:
        payload = source.get(entry.get("source_node", "produtos"), {}).get(entry["legacy_id"], {})
        if not isinstance(payload, dict):
            payload = {"descricao": safe(payload)}
        sku_rows.append([
            entry["legacy_id"], entry.get("legacy_sku"), payload.get("descricao", ""), payload.get("cliente", ""),
            payload.get("clienteKey", ""), payload.get("categoria", ""), payload.get("subcategoria", ""),
            payload.get("volume", ""), payload.get("unidadeVolume", ""), entry.get("classification"), entry.get("reason"),
            "", "", "", "",
        ])
    sku_sheet = add_sheet(workbook, "SKUS_HISTORICOS", [
        "LEGACY_ID", "SKU_LEGADO", "DESCRICAO", "CLIENTE", "CLIENTE_KEY", "CATEGORIA", "SUBCATEGORIA",
        "VOLUME", "UNIDADE_VOLUME", "CLASSIFICACAO", "MOTIVO", "DECISAO", "SKU_ALVO_ID",
        "CODIGO_ATUAL", "OBSERVACAO",
    ], sku_rows)
    add_list_validation(sku_sheet, "L", "ALIAS,MANTER_LEGADO,IGNORAR", len(sku_rows))

    output = Path(args.output)
    output.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output)
    print(json.dumps({
        "status": "CREATED",
        "output": str(output),
        "blocked_suppliers": len(blocked_supplier_rows),
        "materials": len(material_rows),
        "skus": len(sku_rows),
    }, ensure_ascii=False))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
